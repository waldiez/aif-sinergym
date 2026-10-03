#!/usr/bin/env python3
"""
AIF experiment runner for the paper (v4, pairs with aif_agent.py).

Every run writes <out>/<label>/results.json (one record per episode); this
script only reads those files -- no log parsing. Resumable: a run whose JSON
already has the requested number of episodes is skipped.

Each configuration is reported two ways:
  ep1   : the FIRST simulated year from the physics prior (online, unseen year)
  last3 : mean +- std over the last 3 episodes (continued adaptation, same year)

Suites
  core       baseline / learn_C / congestion / full / legacy-v14 reproduction
  sweep      congestion weight {0,1,2,3,5,8} x learn_C {off,on}  -> Pareto front
  horizon    policy_len 4 / 8 / 16 with horizon normalisation
  alt        deadband 2, stochastic, TOU on, coupling (learned/fixed), fixed normaliser
  mechanism  heterogeneous fixed weights; congestion source own/building/lag24/const
  followup   learn_C threshold (c=4), mechanism controls at c=8, 5 more hetero draws
  transfer   A3-A5: every run trained/adapted on the deterministic mixed year (W0)
             is tested on W1 = mixed + noise (--seeds, default 1-5) and
             W2 = hot / cool, two ways:
               frozen : load the W0 checkpoint of each --freeze_from run, no learning
               online : fresh agent from the physics prior (--online_from runs)
             A robustness summary (mean +- std over seeds) is printed per group.
  rbc          RBC reactive / proactive / demand-limiting on W0      (rbc_v2.py)
  rbc_pareto   RBC cool-margin curve + demand-limiting ablations on W0
  rbc_transfer RBC on W1 (same noise seeds as AIF) and W2   (--rbc_from)
  maddpg       MADDPG checkpoints re-scored on W0          (maddpg_eval.py, --maddpg)
  maddpg_transfer  MADDPG checkpoints on W1 / W2
  all        everything above

  python run_ablation_v4.py --suite core --dry          # print commands
  python run_ablation_v4.py --suite core,sweep,horizon,alt,mechanism --jobs 3
  python run_ablation_v4.py --suite transfer \
      --freeze_from 01_baseline,04_full,sw_c8_L1,05_cong2_learnC,mech_const \
      --online_from 01_baseline,04_full --jobs 3 --chdir --clean_eplus
  python run_ablation_v4.py --report                    # tables from existing runs
  python run_ablation_v4.py --only 04_full,mech_const   # specific runs
"""
import argparse, csv, datetime, json, os, shutil, subprocess, sys, threading
from concurrent.futures import ThreadPoolExecutor

HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPT = os.environ.get("AIF_SCRIPT", os.path.join(HERE, "aif_agent.py"))
RBC_SCRIPT = os.environ.get("RBC_SCRIPT", os.path.join(HERE, "rbc_v2.py"))
MADDPG_SCRIPT = os.environ.get("MADDPG_SCRIPT", os.path.join(HERE, "maddpg_eval.py"))
PYTHON = sys.executable

# ---------------------------------------------------------------------------
# Base configuration = the headline controller, with the v15 fixes ON.
# ---------------------------------------------------------------------------
BASE = {
    "--mode": "zone", "--structural": True, "--weather": "mixed",
    "--energy_weight": "0.2", "--deadband_weight": "8.0",
    "--floor_power_idx": "92,93,94", "--floor_power_scale": "0.0011111",
    "--action_temp": "0", "--tou_weight": "0", "--policy_len": "8",
    "--horizon_ref": "8",                       # F4 (identity at policy_len 8)
    "--congestion_weight": "0",
    "--c_lr": "0.02", "--energy_w_min": "0", "--energy_w_max": "1.0",
    "--couple_k": "0.12", "--couple_lr": "0.01",
}
FULL = {"--learn_C": True, "--congestion_weight": "5"}
FIXED_REF_W = "92400"   # incumbent (RBC proactive) annual peak, W -> 'last year's bill'
VAR_SIGMA = "1.5"       # dry-bulb OU sigma (degC) for the stochastic-weather runs


def R(label, ov, episodes=5, depends=None, extra=None, group=None, script="aif"):
    if isinstance(extra, str):
        extra = [extra]
    return {"label": label, "ov": ov, "episodes": episodes, "depends": depends,
            "extra": list(extra or []), "group": group, "script": script}


def RB(label, ov, group=None):
    """An RBC run (deterministic controller: one episode)."""
    return R(label, ov, episodes=1, group=group, script="rbc")


# transfer-suite settings (set from the command line)
TRANSFER_FREEZE = ["01_baseline", "04_full"]
TRANSFER_ONLINE = ["01_baseline", "04_full"]
TRANSFER_SEEDS = [1, 2, 3, 4, 5]
RBC_TRANSFER = ["rbc_reactive", "rbc_proactive", "rbc_dl"]
MADDPG_CKPTS = {}   # name -> {"--checkpoint_dir": ...} or {"--model":..., "--normalizer":...}


def parse_maddpg(spec):
    """'cold=checkpoint-cold,pro=runs/pro/maddpg_office_ep50.pt:runs/pro/maddpg_normalizer.npz'"""
    out = {}
    for item in [x.strip() for x in spec.split(",") if x.strip()]:
        name, path = item.split("=", 1)
        if ":" in path:
            mdl, nrm = path.split(":", 1)
            out[name] = {"--model": os.path.abspath(mdl), "--normalizer": os.path.abspath(nrm)}
        else:
            out[name] = {"--checkpoint_dir": os.path.abspath(path)}
    return out


def suite_core():
    return [
        R("01_baseline", {}),
        R("02_learnC", {"--learn_C": True}),
        R("03_cong5", {"--congestion_weight": "5"}),
        R("04_full", FULL),
        R("05_cong2_learnC", {"--learn_C": True, "--congestion_weight": "2"}),
        # v14 reproduction: must match the old table (sanity + size of the F1 fix)
        R("00_full_legacy_v14", {**FULL, "--legacy_infer": True, "--horizon_ref": "0"}),
    ]


def suite_sweep():
    out = []
    for L in (0, 1):
        for w in ("1", "2", "3", "5", "8"):
            if (L, w) in ((1, "5"), (1, "2"), (0, "5")):
                continue                        # already in core
            ov = {"--congestion_weight": w}
            if L: ov["--learn_C"] = True
            out.append(R(f"sw_c{w}_L{L}", ov))
    return out


def suite_horizon():
    return [R("h04_full", {**FULL, "--policy_len": "4"}),
            R("h16_full", {**FULL, "--policy_len": "16"})]


def suite_alt():
    return [
        R("alt_db2", {**FULL, "--deadband_weight": "2.0"}),
        R("alt_stoch05", {**FULL, "--action_temp": "0.5"}),
        R("alt_tou1", {**FULL, "--tou_weight": "1"}),
        R("alt_couple", {**FULL, "--couple": True}),
        R("alt_couple_fixed", {**FULL, "--couple": True, "--no_couple_learn": True}),
        R("alt_fixedref", {**FULL, "--cong_ref_W": FIXED_REF_W}),
    ]


def suite_mechanism():
    runs = []
    for s in (0, 1, 2):   # heterogeneous FIXED weights + congestion, no learn_C
        runs.append(R(f"mech_hetero_s{s}", {"--congestion_weight": "5",
                      "--energy_w_init": "uniform", "--energy_w_seed": str(s)}))
    runs.append(R("mech_hetero_nocong_s0", {"--energy_w_init": "uniform",
                                            "--energy_w_seed": "0"}))
    for src in ("own", "building", "lag24"):
        runs.append(R(f"mech_{src}", {**FULL, "--cong_source": src}))
    # const: same mean signal as 04_full (per floor, its last episode), no dynamics
    runs.append(R("mech_const", {**FULL, "--cong_source": "const"},
                  depends="04_full", extra="cong_const_from_dep"))
    return runs


def _w0_runs():
    """All W0 runs (the non-transfer suites), by label."""
    return {r["label"]: r for n in ORDER
            if "transfer" not in n and not n.startswith(("rbc", "maddpg"))
            for r in SUITES[n]()}


def _test_conditions():
    conds = [(f"var_s{s}", {"--weather": "mixed", "--weather_variability": VAR_SIGMA,
                            "--seed": str(s)}, "W1") for s in TRANSFER_SEEDS]
    conds += [(w, {"--weather": w}, "W2_" + w) for w in ("hot", "cool")]
    return conds


def suite_transfer():
    w0 = _w0_runs()
    runs = []
    for mode, sources in (("frz", TRANSFER_FREEZE), ("onl", TRANSFER_ONLINE)):
        for src in sources:
            if src not in w0:
                raise SystemExit(f"--{'freeze' if mode == 'frz' else 'online'}_from: "
                                 f"unknown run '{src}'")
            base = w0[src]
            ov = dict(base["ov"])
            extra = []
            # a source that used a measured constant congestion signal keeps it
            if "cong_const_from_dep" in base["extra"]:
                extra.append("cong_const_from_src")
            if mode == "frz":
                ov["--freeze_all"] = True
                extra.append("load_dep_checkpoint")
            for tag, wov, gname in _test_conditions():
                # online runs only wait for the source when they need its data
                runs.append(R(f"tr_{mode}_{src}_{tag}", {**ov, **wov}, episodes=1,
                              depends=(src if extra else None), extra=extra,
                              group=f"{mode}|{src}|{gname}"))
    return runs


def suite_followup():
    """Follow-ups from the A2 results.
    (i)  locate the learn_C threshold between congestion 3 and 5;
    (ii) the four mechanism controls at congestion 8: do they fail because the
         signal lacks the right information, or only because they sit just
         below the regime threshold?
    (iii) five more heterogeneous-weight draws: how often does RANDOM
         heterogeneity reach the low-peak regime (1 of 3 so far)?"""
    runs = [R("sw_c4_L1", {"--learn_C": True, "--congestion_weight": "4"})]
    for src in ("own", "building", "lag24"):
        runs.append(R(f"mech_{src}_c8", {**FULL, "--congestion_weight": "8",
                                         "--cong_source": src}))
    runs.append(R("mech_const_c8", {**FULL, "--congestion_weight": "8",
                                    "--cong_source": "const"},
                  depends="04_full", extra="cong_const_from_dep"))
    for sd in (3, 4, 5, 6, 7):
        runs.append(R(f"mech_hetero_s{sd}", {"--congestion_weight": "5",
                      "--energy_w_init": "uniform", "--energy_w_seed": str(sd)}))
    return runs


def suite_rbc():
    """The three rule-based baselines on W0 (deterministic mixed year)."""
    return [RB("rbc_reactive", {"--variant": "reactive"}),
            RB("rbc_proactive", {"--variant": "proactive"}),
            RB("rbc_dl", {"--variant": "dl"})]


def suite_rbc_pareto():
    """RBC comfort-peak trade-off curve on W0, plus demand-limiting ablations."""
    runs = [RB(f"rbc_pro_cm{str(c).replace('.', 'p')}",
               {"--variant": "proactive", "--cool_margin": str(c)})
            for c in (0.5, 1.5, 2.0, 2.5)]
    runs += [RB(f"rbc_dl_f{str(f).replace('.', 'p')}",
                {"--variant": "dl", "--dl_limit_frac": str(f)}) for f in (0.95, 0.85, 0.8)]
    runs += [RB("rbc_dl_rotate", {"--variant": "dl", "--dl_mode": "rotate"}),
             RB("rbc_dl_shed2", {"--variant": "dl", "--dl_shed_C": "2.0"}),
             RB("rbc_stagger_only", {"--variant": "dl", "--dl_limit_frac": "0"}),
             RB("rbc_limit_only", {"--variant": "dl", "--stagger_min": "0"})]
    return runs


def suite_rbc_transfer():
    """RBC on W1 (the SAME noise seeds as the AIF runs) and W2."""
    w0 = {r["label"]: r for n in ("rbc", "rbc_pareto") for r in SUITES[n]()}
    runs = []
    for src in RBC_TRANSFER:
        if src not in w0:
            raise SystemExit(f"--rbc_from: unknown RBC run '{src}'")
        for tag, wov, gname in _test_conditions():
            short = src[4:] if src.startswith("rbc_") else src
            runs.append(RB(f"tr_rbc_{short}_{tag}", {**w0[src]["ov"], **wov},
                           group=f"rbc|{src}|{gname}"))
    return runs


def suite_maddpg():
    """Existing MADDPG checkpoints re-scored on W0 (deterministic actors)."""
    return [R(f"maddpg_{n}", dict(ov), episodes=1, script="maddpg")
            for n, ov in MADDPG_CKPTS.items()]


def suite_maddpg_transfer():
    """MADDPG checkpoints (trained on W0) on W1 (same noise seeds) and W2."""
    runs = []
    for n, ov in MADDPG_CKPTS.items():
        for tag, wov, gname in _test_conditions():
            runs.append(R(f"tr_mad_{n}_{tag}", {**ov, **wov}, episodes=1, script="maddpg",
                          group=f"mad|maddpg_{n}|{gname}"))
    return runs


SUITES = {"core": suite_core, "sweep": suite_sweep, "horizon": suite_horizon,
          "alt": suite_alt, "mechanism": suite_mechanism, "followup": suite_followup,
          "transfer": suite_transfer, "rbc": suite_rbc, "rbc_pareto": suite_rbc_pareto,
          "rbc_transfer": suite_rbc_transfer, "maddpg": suite_maddpg,
          "maddpg_transfer": suite_maddpg_transfer}
ORDER = ["core", "sweep", "horizon", "alt", "mechanism", "followup", "transfer",
         "rbc", "rbc_pareto", "rbc_transfer", "maddpg", "maddpg_transfer"]

# ---------------------------------------------------------------------------
METRICS = ["peak_kW", "CF", "LF", "CS", "CS_occ", "ZCR", "mean_dev", "db_viol",
           "energy_kWh", "en_cost", "dem_cost", "tot_cost", "sum_monthly_peaks_kW",
           "override_rate", "custom_reward", "linear_reward", "hourly_reward",
           "core_k", "perim_k", "prederr_impr"]


def out_dir(root, label): return os.path.join(root, label)
def res_path(root, label): return os.path.join(out_dir(root, label), "results.json")


def load_res(root, label):
    p = res_path(root, label)
    if not os.path.exists(p):
        return None
    try:
        return json.load(open(p))
    except Exception:
        return None


def build_cmd(run, root):
    if run.get("script") in ("rbc", "maddpg"):
        a = dict(run["ov"])
        if run["script"] == "maddpg" and BASE.get("--device"):
            a["--device"] = BASE["--device"]
        a["--out_dir"] = os.path.abspath(out_dir(root, run["label"]))
        a["--label"] = run["label"]
        cmd = [PYTHON, RBC_SCRIPT if run["script"] == "rbc" else MADDPG_SCRIPT]
        for k, v in a.items():
            if v is True: cmd.append(k)
            elif v is False or v is None: pass
            else: cmd += [k, str(v)]
        return cmd
    a = dict(BASE); a.update(run["ov"]); a["--episodes"] = str(run["episodes"])
    od = os.path.abspath(out_dir(root, run["label"]))
    a["--out_dir"] = od; a["--label"] = run["label"]
    a["--checkpoint"] = os.path.join(od, "checkpoint.pkl")
    if "load_dep_checkpoint" in run["extra"]:
        a["--load"] = os.path.abspath(os.path.join(out_dir(root, run["depends"]),
                                                   "checkpoint.pkl"))
    if "cong_const_from_src" in run["extra"]:
        src = load_res(root, run["depends"])
        cc = (src or {}).get("config", {}).get("cong_const")
        if not cc:
            raise RuntimeError(f"{run['label']} needs finished {run['depends']}")
        a["--cong_const"] = cc
    if "cong_const_from_dep" in run["extra"]:
        dep = load_res(root, run["depends"])
        if not dep or not dep["episodes"]:
            raise RuntimeError(f"{run['label']} needs finished {run['depends']}")
        cc = dep["episodes"][-1]["cong_obs_mean_by_floor"]
        a["--cong_const"] = ",".join(f"{x:.5f}" for x in cc)
    cmd = [PYTHON, SCRIPT]
    for k, v in a.items():
        if v is True: cmd.append(k)
        elif v is False or v is None: pass
        else: cmd += [k, str(v)]
    return cmd


def is_done(run, root):
    r = load_res(root, run["label"])
    return bool(r) and len(r.get("episodes", [])) >= run["episodes"]


_print_lock = threading.Lock()


def execute(run, root, chdir, clean):
    od = out_dir(root, run["label"]); os.makedirs(od, exist_ok=True)
    if is_done(run, root):
        with _print_lock: print(f"[skip] {run['label']} (done)")
        return True
    # a partial JSON from an interrupted run must not be mixed with a new one
    if os.path.exists(res_path(root, run["label"])):
        os.remove(res_path(root, run["label"]))
    cmd = build_cmd(run, root)
    with _print_lock:
        print(f"[{datetime.datetime.now():%H:%M}] run  {run['label']}", flush=True)
    t0 = datetime.datetime.now()
    with open(os.path.join(od, "run.log"), "w") as lf:
        lf.write(" ".join(cmd) + "\n\n"); lf.flush()
        p = subprocess.run(cmd, stdout=lf, stderr=subprocess.STDOUT,
                           cwd=(os.path.abspath(od) if chdir else None))
    if clean and chdir:
        for d in os.listdir(od):
            if d.startswith("eplus-env-") or d.startswith("Eplus-env-"):
                shutil.rmtree(os.path.join(od, d), ignore_errors=True)
    ok = p.returncode == 0 and is_done(run, root)
    with _print_lock:
        dt = (datetime.datetime.now() - t0).total_seconds() / 60
        print(f"[{datetime.datetime.now():%H:%M}] {'done' if ok else 'FAIL'} "
              f"{run['label']} ({dt:.0f} min){'' if ok else ' -> see run.log'}",
              flush=True)
    return ok


EPISODE_CAP = None   # --episodes override (smoke tests), applied everywhere


def _cap(runs):
    if EPISODE_CAP:
        for r in runs: r["episodes"] = min(r["episodes"], EPISODE_CAP)
    return runs


def _all_runs():
    return {r["label"]: r for r in _cap([r for s in ORDER for r in SUITES[s]()])}


def dep_ok(run, root):
    """A dependency is satisfied when its run finished all its episodes (and,
    for warm starts, its checkpoint exists)."""
    if not run["depends"]:
        return True
    dep = _all_runs()[run["depends"]]
    if not is_done(dep, root):
        return False
    if "load_dep_checkpoint" in run["extra"]:
        return os.path.exists(os.path.join(out_dir(root, dep["label"]), "checkpoint.pkl"))
    return True


def run_all(runs, root, jobs, chdir, clean, dry):
    if dry:
        for r in runs:
            try: print(r["label"], ":", " ".join(build_cmd(r, root)))
            except RuntimeError: print(r["label"], f": (runs after {r['depends']})")
        return
    # dependency-ordered waves; each wave runs in parallel
    pending = list(runs)
    while pending:
        ready = [r for r in pending if dep_ok(r, root)]
        if not ready:
            for r in pending:
                print(f"[blocked] {r['label']}: dependency {r['depends']} not finished "
                      f"(select it too, or check its run.log)")
            break
        with ThreadPoolExecutor(max_workers=max(1, jobs)) as ex:
            list(ex.map(lambda r: execute(r, root, chdir, clean), ready))
        pending = [r for r in pending if r not in ready]


# ---------------------------------------------------------------------------
# Reporting
# ---------------------------------------------------------------------------
def summarise(res, last_n=3):
    eps = res["episodes"]
    row = {}
    e1 = eps[0]
    for m in METRICS:
        row[f"{m}@ep1"] = e1.get(m)
    tail = eps[-last_n:] if len(eps) >= last_n else eps
    for m in METRICS:
        vals = [e.get(m) for e in tail if e.get(m) is not None]
        if vals:
            mu = sum(vals) / len(vals)
            sd = (sum((v - mu) ** 2 for v in vals) / max(len(vals) - 1, 1)) ** 0.5
            row[f"{m}@last"] = mu; row[f"{m}@last_std"] = sd
    pe1 = eps[0].get("peak_event") or {}
    row["peak_when_ep1"] = (f"{pe1.get('month')}/{pe1.get('day')} h{pe1.get('hour')}"
                            if pe1 else "")
    pe = eps[-1].get("peak_event") or {}
    row["peak_when"] = f"{pe.get('month')}/{pe.get('day')} h{pe.get('hour')}" if pe else ""
    row["peak_heat/cool/db"] = (f"{pe.get('n_heating')}/{pe.get('n_cooling')}/"
                                f"{pe.get('n_band_lt2')}") if pe else ""
    row["n_eps"] = len(eps)
    row["started_fresh"] = res.get("started_fresh")
    row["metrics_md5"] = res.get("metrics_utils_md5")
    return row


def report(root, runs, path_csv, path_xlsx):
    rows = []
    for r in runs:
        res = load_res(root, r["label"])
        if not res or not res.get("episodes"):
            continue
        row = {"label": r["label"], **summarise(res)}
        row["complete"] = row["n_eps"] >= r["episodes"]
        row["group"] = r.get("group")
        rows.append(row)
    if not rows:
        print("no finished runs"); return
    cols = ["label", "complete", "n_eps", "group", "started_fresh", "metrics_md5",
            "peak_when_ep1", "peak_when", "peak_heat/cool/db"]
    for m in METRICS:
        cols += [f"{m}@ep1", f"{m}@last", f"{m}@last_std"]
    with open(path_csv, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore"); w.writeheader()
        for row in rows: w.writerow(row)
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "AIF runs"; ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="404040")
        for row in rows: ws.append([row.get(c) for c in cols])
        ws.freeze_panes = "B2"; wb.save(path_xlsx)
    except ImportError:
        pass
    md5s = {r.get("metrics_md5") for r in rows}
    if len(md5s) > 1:
        print(f"WARNING: runs were scored with different metrics_utils versions: {md5s}")
    # console table
    hdr = f"{'label':26s} {'peak1':>6s} {'peakL':>6s} {'CF1':>6s} {'CFL':>6s} " \
          f"{'CS1':>6s} {'CSL':>6s} {'ZCRL':>5s} {'$L':>7s} {'ovr%':>5s}  peak@"
    print(hdr); print("-" * len(hdr))
    f = lambda v, fmt: (format(v, fmt) if isinstance(v, (int, float)) else "-")
    for row in rows:
        lab = row["label"] + ("" if row["complete"] else " *")
        print(f"{lab[:26]:26s} {f(row.get('peak_kW@ep1'),'6.1f')} "
              f"{f(row.get('peak_kW@last'),'6.1f')} {f(row.get('CF@ep1'),'6.3f')} "
              f"{f(row.get('CF@last'),'6.3f')} {f(row.get('CS@ep1'),'6.3f')} "
              f"{f(row.get('CS@last'),'6.3f')} {f(row.get('ZCR@last'),'5.1f')} "
              f"{f(row.get('tot_cost@last'),'7.0f')} "
              f"{f(100*row['override_rate@last'] if row.get('override_rate@last') is not None else None,'5.2f')}"
              f"  {row['peak_when']}")
    if not all(r["complete"] for r in rows):
        print("  * = still running / incomplete: '@last' averages only the episodes so far")
    print(f"\n-> {path_csv}" + (f", {path_xlsx}" if os.path.exists(path_xlsx) else ""))
    robustness_summary([r for r in rows if r.get("group") and r["complete"]],
                       os.path.join(os.path.dirname(path_csv), "robustness.csv"))
    weather_check(os.path.dirname(path_csv))


def weather_check(root):
    """Paired comparisons need identical weather: compare the outdoor-temperature
    column of every transfer run with the first run of the same condition."""
    import glob
    ref, bad, n = {}, [], 0
    for d in sorted(glob.glob(os.path.join(root, "tr_*"))):
        tag = d.rsplit("_", 1)[-1] if "_var_s" not in d else "var_s" + d.rsplit("_var_s", 1)[-1]
        z = os.path.join(d, "zone_temps_actions.csv")
        if not os.path.exists(z):
            continue
        with open(z) as fh:
            rd = csv.reader(fh); hdr = next(rd); j = hdr.index("outdoor_temp")
            col = [row[j] for row, _ in zip(rd, range(35040)) if row[0] == "1"]
        n += 1
        if tag not in ref:
            ref[tag] = (os.path.basename(d), col)
        elif col != ref[tag][1]:
            bad.append((os.path.basename(d), ref[tag][0]))
    if n:
        print(f"\nweather check: {n} transfer runs, {len(ref)} conditions -> "
              + ("all identical within condition" if not bad else f"{len(bad)} MISMATCHES"))
        for a_, b_ in bad[:10]:
            print(f"   {a_} differs from {b_}")


def robustness_summary(rows, path):
    """Transfer runs: one line per (mode, source, condition) group; W1 groups
    aggregate the noise seeds (mean +- sample std of the single test episode)."""
    if not rows:
        return
    groups = {}
    for r in rows:
        groups.setdefault(r["group"], []).append(r)
    keys = ["peak_kW", "CF", "CS", "CS_occ", "ZCR", "tot_cost", "energy_kWh", "override_rate"]
    out = []
    for g, rs in sorted(groups.items()):
        line = {"group": g, "n": len(rs)}
        for k in keys:
            v = [x[f"{k}@ep1"] for x in rs if x.get(f"{k}@ep1") is not None]
            if v:
                mu = sum(v) / len(v)
                sd = (sum((y - mu) ** 2 for y in v) / (len(v) - 1)) ** 0.5 if len(v) > 1 else 0.0
                line[k], line[k + "_std"] = mu, sd
        out.append(line)
    with open(path, "w", newline="") as fh:
        cols = ["group", "n"] + [c for k in keys for c in (k, k + "_std")]
        w = csv.DictWriter(fh, fieldnames=cols, extrasaction="ignore"); w.writeheader()
        for line in out: w.writerow(line)
    print("\nROBUSTNESS / GENERALISATION (single test year per run; W1 = mean +- std over seeds)")
    hdr = f"{'mode|source|condition':34s} {'n':>2s} {'peak kW':>12s} {'CF':>13s} " \
          f"{'CS':>13s} {'ZCR %':>11s} {'cost $':>14s}"
    print(hdr); print("-" * len(hdr))
    pm = lambda l, k, fmt: (f"{format(l[k], fmt)}±{format(l[k+'_std'], fmt)}"
                            if k in l else "-")
    for l in out:
        print(f"{l['group'][:34]:34s} {l['n']:>2d} {pm(l,'peak_kW','.1f'):>12s} "
              f"{pm(l,'CF','.3f'):>13s} {pm(l,'CS','.3f'):>13s} {pm(l,'ZCR','.1f'):>11s} "
              f"{pm(l,'tot_cost','.0f'):>14s}")
    print(f"-> {path}")


def main():
    global TRANSFER_FREEZE, TRANSFER_ONLINE, TRANSFER_SEEDS, EPISODE_CAP, RBC_TRANSFER, \
        MADDPG_CKPTS
    ap = argparse.ArgumentParser()
    ap.add_argument("--suite", default="core",
                    help="comma list of " + ",".join(ORDER) + " or 'all'")
    ap.add_argument("--only", default=None, help="comma list of labels")
    ap.add_argument("--root", default="ab_runs")
    ap.add_argument("--jobs", type=int, default=1)
    ap.add_argument("--chdir", action="store_true",
                    help="run each job inside its own out_dir (needed for --jobs>1 so "
                         "Sinergym working dirs do not collide)")
    ap.add_argument("--clean_eplus", action="store_true",
                    help="delete Sinergym/EnergyPlus output dirs after each run (needs --chdir)")
    ap.add_argument("--episodes", type=int, default=None, help="override episodes (smoke tests)")
    ap.add_argument("--device", default=None, help="auto | cuda | cpu (passed through)")
    ap.add_argument("--freeze_from", default=",".join(TRANSFER_FREEZE),
                    help="transfer: W0 runs whose checkpoint is tested frozen")
    ap.add_argument("--online_from", default=",".join(TRANSFER_ONLINE),
                    help="transfer: W0 configs re-run fresh (online) on the test weather")
    ap.add_argument("--seeds", default=",".join(map(str, TRANSFER_SEEDS)),
                    help="transfer: noise seeds for W1")
    ap.add_argument("--maddpg", default="",
                    help="maddpg/maddpg_transfer: name=checkpoint_dir or "
                         "name=model.pt:normalizer.npz, comma-separated")
    ap.add_argument("--rbc_from", default=",".join(RBC_TRANSFER),
                    help="rbc_transfer: RBC runs to test on W1/W2")
    ap.add_argument("--dry", action="store_true")
    ap.add_argument("--report", action="store_true", help="only write the tables")
    a = ap.parse_args()

    TRANSFER_FREEZE = [x.strip() for x in a.freeze_from.split(",") if x.strip()]
    TRANSFER_ONLINE = [x.strip() for x in a.online_from.split(",") if x.strip()]
    TRANSFER_SEEDS = [int(x) for x in a.seeds.split(",") if x.strip()]
    RBC_TRANSFER = [x.strip() for x in a.rbc_from.split(",") if x.strip()]
    MADDPG_CKPTS = parse_maddpg(a.maddpg)
    names = ORDER if a.suite == "all" else [s.strip() for s in a.suite.split(",")]
    runs = [r for s in names for r in SUITES[s]()]
    if a.only:
        want = {x.strip() for x in a.only.split(",")}
        allr = [r for s in ORDER for r in SUITES[s]()]
        runs = [r for r in allr if r["label"] in want]
    EPISODE_CAP = a.episodes
    _cap(runs)
    if a.device:
        BASE["--device"] = a.device
    if a.jobs > 1 and not a.chdir:
        print("NOTE: --jobs > 1 without --chdir: Sinergym working dirs may collide; "
              "consider --chdir.")
    os.makedirs(a.root, exist_ok=True)
    if not a.report:
        run_all(runs, a.root, a.jobs, a.chdir, a.clean_eplus, a.dry)
    if not a.dry:
        report(a.root, runs, os.path.join(a.root, "results.csv"),
               os.path.join(a.root, "results.xlsx"))


if __name__ == "__main__":
    main()
